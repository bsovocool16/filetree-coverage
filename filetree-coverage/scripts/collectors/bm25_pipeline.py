#!/usr/bin/env python3
"""
coverage_pipeline.py — retrieval over a folder of documents that records, for
every file in the folder and every question, which of these happened:

    unindexed   in the folder, but no text could be extracted
                (scanned PDF with no text layer, unsupported type, empty file)
    untouched   indexed, but no chunk of it scored into the candidate set
    retrieved   at least one chunk was a retrieval candidate (top-K)
    in_context  at least one chunk was actually placed in the model's prompt
    cited       the model cited it in the answer

Output is one coverage.json that render_viewer.py turns into the coverage map.

Three ways to run
-----------------
1. Claude (or any reader) answers from the excerpts — no API key:

    coverage_pipeline.py retrieve <folder> -q "..." [-q "..."] -o <run>/coverage.json
        writes coverage.json (tiers through in_context) and <run>/<qid>.excerpts.md
    ... the model reads <qid>.excerpts.md and writes <qid>.answer.md citing [[n]] ...
    coverage_pipeline.py record <run>/coverage.json --q q1 --answer <run>/q1.answer.md
        parses the [[n]] citations into the cited tier

2. One shot, no model (placeholder answer that cites the strongest excerpts):

    coverage_pipeline.py run <folder> -q "..." -o coverage.json --generator stub

3. One shot through the Anthropic API (pip install anthropic; ANTHROPIC_API_KEY):

    coverage_pipeline.py run <folder> -q "..." -o coverage.json --generator claude

Expected files (the gap list) can be declared with --expect GLOB (repeatable,
matched against the file's path inside the folder) and refined in the viewer.

Dependencies: standard library only. Used if installed: pdfplumber or pypdf
(PDF), python-docx (.docx), openpyxl (.xlsx); `pdftotext` on PATH is a PDF fallback.
"""

from __future__ import annotations

import argparse
import datetime as _dt
import fnmatch
import json
import math
import re
import shutil
import subprocess
import sys
from collections import Counter
from dataclasses import dataclass, asdict
from pathlib import Path

# ----------------------------------------------------------------------------
# Text extraction
# ----------------------------------------------------------------------------

TEXT_EXT = {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".rtf", ".log"}
HTML_EXT = {".html", ".htm", ".xml"}
SKIP_NAMES = {".DS_Store", "Thumbs.db"}


def _read_text(path: Path) -> str:
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return ""


def _extract_pdf(path: Path) -> str:
    try:
        import pdfplumber  # type: ignore
        with pdfplumber.open(str(path)) as pdf:
            return "\n\n".join((p.extract_text() or "") for p in pdf.pages)
    except Exception:
        pass
    try:
        from pypdf import PdfReader  # type: ignore
        return "\n\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
    except Exception:
        pass
    if shutil.which("pdftotext"):
        try:
            return subprocess.run(["pdftotext", "-layout", str(path), "-"],
                                  capture_output=True, text=True, timeout=120).stdout
        except Exception:
            pass
    return ""


def _extract_docx(path: Path) -> str:
    try:
        import docx  # type: ignore
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append("\t".join(c.text for c in row.cells))
        return "\n".join(parts)
    except Exception:
        return ""


def _extract_xlsx(path: Path) -> str:
    try:
        import openpyxl  # type: ignore
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"# Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append("\t".join(cells))
        return "\n".join(parts)
    except Exception:
        return ""


def extract_text(path: Path) -> tuple[str, str | None]:
    """Return (text, failure_reason); failure_reason is None on success."""
    ext = path.suffix.lower()
    if ext in TEXT_EXT:
        text = _read_text(path)
    elif ext in HTML_EXT:
        text = re.sub(r"<[^>]+>", " ", _read_text(path))
    elif ext == ".pdf":
        text = _extract_pdf(path)
        if not text.strip():
            return "", "PDF has no text layer (scanned image?) or no PDF library installed"
    elif ext == ".docx":
        text = _extract_docx(path)
        if not text.strip():
            return "", "could not read .docx (install python-docx)"
    elif ext == ".xlsx":
        text = _extract_xlsx(path)
        if not text.strip():
            return "", "could not read .xlsx (install openpyxl)"
    else:
        return "", f"unsupported file type {ext or '(none)'}"
    if not text.strip():
        return "", "file is empty"
    return text, None


# ----------------------------------------------------------------------------
# Chunking
# ----------------------------------------------------------------------------

def chunk_text(text: str, size: int = 1200, overlap: int = 150) -> list[tuple[int, int, str]]:
    text = text.replace("\r\n", "\n")
    n = len(text)
    out: list[tuple[int, int, str]] = []
    start = 0
    while start < n:
        end = min(n, start + size)
        if end < n:
            window = text[start + int(size * 0.7): end]
            cut = max(window.rfind("\n\n"), window.rfind(". "))
            if cut > 0:
                end = start + int(size * 0.7) + cut + 1
        piece = text[start:end]
        if piece.strip():
            out.append((start, end, piece))
        if end >= n:
            break
        start = max(end - overlap, start + 1)
    return out


# ----------------------------------------------------------------------------
# BM25 (self-contained)
# ----------------------------------------------------------------------------

STOP = set("""a an and are as at be by for from has have in is it its of on or that the this to was were will with
which who whom whose what when where how any all each such than then there these those into under over shall may
not no nor if but do does did been being also per via""".split())

_tok = re.compile(r"[a-z0-9]+")

_SUFFIXES = (
    ("ational", "ate"), ("tional", "tion"), ("ization", "ize"), ("ations", "ate"), ("ation", "ate"),
    ("ments", ""), ("ment", ""), ("ness", ""), ("ities", "ity"), ("ity", ""),
    ("ingly", ""), ("ings", ""), ("ing", ""), ("edly", ""), ("ed", ""), ("ies", "y"), ("sses", "ss"),
    ("es", ""), ("s", ""), ("ance", ""), ("ence", ""), ("able", ""), ("ible", ""), ("ive", ""), ("ize", ""),
    ("al", ""), ("er", ""), ("ly", ""),
)


def stem(t: str) -> str:
    if len(t) <= 3:
        return t
    for suf, rep in _SUFFIXES:
        if t.endswith(suf) and len(t) - len(suf) + len(rep) >= 3:
            t = t[: -len(suf)] + rep
            break
    if len(t) > 4 and t.endswith("e"):
        t = t[:-1]
    return t


def tokenize(s: str) -> list[str]:
    return [stem(t) for t in _tok.findall(s.lower()) if t not in STOP and len(t) >= 2]


class BM25:
    def __init__(self, docs: list[list[str]], k1: float = 1.5, b: float = 0.75):
        self.k1, self.b, self.docs, self.N = k1, b, docs, len(docs)
        self.avgdl = (sum(len(d) for d in docs) / self.N) if self.N else 0.0
        self.tf = [Counter(d) for d in docs]
        df: Counter = Counter()
        for tf in self.tf:
            df.update(tf.keys())
        self.idf = {t: math.log(1 + (self.N - n + 0.5) / (n + 0.5)) for t, n in df.items()}

    def score(self, query: list[str]) -> list[float]:
        scores = [0.0] * self.N
        for term in set(query):
            idf = self.idf.get(term)
            if idf is None:
                continue
            for i, tf in enumerate(self.tf):
                f = tf.get(term)
                if not f:
                    continue
                dl = len(self.docs[i])
                scores[i] += idf * f * (self.k1 + 1) / (f + self.k1 * (1 - self.b + self.b * dl / self.avgdl))
        return scores


# ----------------------------------------------------------------------------
# Index
# ----------------------------------------------------------------------------

@dataclass
class FileRec:
    path: str
    size: int
    indexed: bool
    chunks: int
    reason: str | None = None


@dataclass
class Chunk:
    file: str
    n: int
    start: int
    end: int
    text: str


def build_index(root: Path, chunk_chars: int, overlap: int):
    files: list[FileRec] = []
    chunks: list[Chunk] = []
    paths = sorted(p for p in root.rglob("*") if p.is_file() and p.name not in SKIP_NAMES
                   and not any(part.startswith(".") for part in p.relative_to(root).parts))
    for p in paths:
        rel = p.relative_to(root).as_posix()
        text, reason = extract_text(p)
        if reason:
            files.append(FileRec(rel, p.stat().st_size, False, 0, reason))
            print(f"  [unindexed] {rel}: {reason}", file=sys.stderr)
            continue
        pieces = chunk_text(text, chunk_chars, overlap)
        for i, (s, e, t) in enumerate(pieces):
            chunks.append(Chunk(rel, i, s, e, t))
        files.append(FileRec(rel, p.stat().st_size, True, len(pieces)))
    print(f"Indexed {sum(1 for f in files if f.indexed)}/{len(files)} files, {len(chunks)} chunks", file=sys.stderr)
    return files, chunks


# ----------------------------------------------------------------------------
# Retrieval + context assembly
# ----------------------------------------------------------------------------

def retrieve(bm25: BM25, question: str, k: int):
    scores = bm25.score(tokenize(question))
    ranked = sorted(((s, i) for i, s in enumerate(scores) if s > 0), reverse=True)
    return [(i, s) for s, i in ranked[:k]]


def assemble_context(cands, chunks: list[Chunk], m: int, per_file_cap: int):
    """The chunks that actually go to the model: top-M by score with a per-file
    cap so one long document cannot crowd out the rest."""
    chosen, per_file = [], Counter()
    for i, s in cands:
        f = chunks[i].file
        if per_file[f] >= per_file_cap:
            continue
        chosen.append((i, s))
        per_file[f] += 1
        if len(chosen) >= m:
            break
    return chosen


# ----------------------------------------------------------------------------
# Generation
# ----------------------------------------------------------------------------

SYSTEM_PROMPT = """You are assisting with a document review over a folder of files.
Answer the question using ONLY the numbered excerpts provided. Every factual
statement must cite the excerpt(s) it rests on using the form [[n]] where n is
the excerpt number. If the excerpts do not answer the question, say so
explicitly and cite nothing. Do not speculate about documents you were not shown."""

CITE_RE = re.compile(r"\[\[(\d+)\]\]")


def excerpts_markdown(question: str, excerpts: list[dict], qid: str) -> str:
    lines = [f"# {qid}: {question}", "",
             "Answer from these excerpts only. Cite each statement to its excerpt as [[n]]. "
             "If the excerpts do not answer the question, say so and cite nothing. "
             "Do not speculate about documents you were not shown.", ""]
    for e in excerpts:
        lines.append(f"## [[{e['n']}]] {e['file']} — chunk {e['chunk'] + 1} of {e['file_chunks']} (score {e['score']})")
        lines.append("")
        lines.append(e["text"].strip())
        lines.append("")
    return "\n".join(lines)


def generate_claude(question: str, excerpts: list[dict], model: str) -> str:
    try:
        import anthropic  # type: ignore
    except ImportError:
        sys.exit("--generator claude needs the anthropic package: pip install anthropic")
    body = "\n\n".join(f"[[{e['n']}]] ({e['file']}, chunk {e['chunk'] + 1} of {e['file_chunks']})\n{e['text']}"
                       for e in excerpts)
    msg = anthropic.Anthropic().messages.create(
        model=model, max_tokens=1500, system=SYSTEM_PROMPT,
        messages=[{"role": "user", "content": f"QUESTION: {question}\n\nEXCERPTS:\n\n{body}"}])
    return "".join(getattr(b, "text", "") for b in msg.content)


def generate_stub(question: str, excerpts: list[dict]) -> str:
    """No model call: cites the files whose best excerpt scored within 60% of the
    top score, so the cited tier exercises the viewer. Labelled as a stub."""
    if not excerpts:
        return "STUB (no model call): no excerpts were retrieved for this question."
    best = max(e["score"] for e in excerpts)
    by_file: dict[str, dict] = {}
    for e in excerpts:
        if e["file"] not in by_file or e["score"] > by_file[e["file"]]["score"]:
            by_file[e["file"]] = e
    lines = ["STUB (no model call): the strongest excerpts for this question were:"]
    for e in sorted((e for e in by_file.values() if e["score"] >= 0.6 * best), key=lambda e: -e["score"]):
        lines.append(f"- {e['file']} [[{e['n']}]]")
    return "\n".join(lines)


# ----------------------------------------------------------------------------
# Tiers
# ----------------------------------------------------------------------------

def tier_files(cands, ctx, chunks, file_chunk_count):
    per_file: dict[str, dict] = {}
    for rank, (i, s) in enumerate(cands, 1):
        c = chunks[i]
        d = per_file.setdefault(c.file, {
            "tier": "retrieved", "chunks_retrieved": 0, "chunks_in_context": 0, "best_score": 0.0,
            "best_rank": rank, "chunks_total": file_chunk_count[c.file], "context_chunks": []})
        d["chunks_retrieved"] += 1
        d["best_score"] = max(d["best_score"], round(s, 3))
    for i, s in ctx:
        c = chunks[i]
        d = per_file[c.file]
        d["chunks_in_context"] += 1
        d["context_chunks"].append(c.n + 1)
        d["tier"] = "in_context"
    return per_file


def apply_citations(q: dict, answer: str) -> set[str]:
    cited_n = {int(m) for m in CITE_RE.findall(answer)}
    cited_files = {e["file"] for e in q["excerpts"] if e["n"] in cited_n}
    for f, d in q["files"].items():
        if f in cited_files:
            d["tier"] = "cited"
        elif d["tier"] == "cited":
            d["tier"] = "in_context"
    q["answer"] = answer
    q["summary"]["files_cited"] = len(cited_files)
    return cited_files


def summarize(files, per_file, cands, ctx):
    return {
        "files_total": len(files),
        "files_indexed": sum(1 for f in files if f.indexed),
        "files_retrieved": len(per_file),
        "files_in_context": sum(1 for d in per_file.values() if d["tier"] in ("in_context", "cited")),
        "files_cited": sum(1 for d in per_file.values() if d["tier"] == "cited"),
        "chunks_retrieved": len(cands),
        "chunks_in_context": len(ctx),
    }


def expected_from_globs(files, globs):
    if not globs:
        return []
    return sorted({f.path for f in files for g in globs
                   if fnmatch.fnmatch(f.path, g) or fnmatch.fnmatch(f.path, g.rstrip("/") + "/*")})


# ----------------------------------------------------------------------------
# Commands
# ----------------------------------------------------------------------------

def _questions(args):
    qs = list(args.q or [])
    if args.questions:
        qs += [l.strip() for l in Path(args.questions).read_text().splitlines() if l.strip()]
    if not qs:
        sys.exit("give at least one -q question or --questions file")
    return qs


def _prepare(args, generator_label: str):
    root = Path(args.folder).resolve()
    if not root.is_dir():
        sys.exit(f"not a directory: {root}")
    files, chunks = build_index(root, args.chunk, args.overlap)
    bm25 = BM25([tokenize(c.text) for c in chunks])
    file_chunk_count = {f.path: f.chunks for f in files}
    out = {
        "generated_at": _dt.datetime.now().isoformat(timespec="seconds"),
        "vdr_root": root.name,
        "folder": str(root),
        "pipeline": {"retriever": "bm25", "k_retrieved": args.k, "k_context": args.context,
                     "per_file_cap": args.per_file_cap, "chunk_chars": args.chunk, "generator": generator_label},
        "files": [asdict(f) for f in files],
        "questions": [],
    }
    expect_map = {}
    if args.expect_file:
        expect_map = json.loads(Path(args.expect_file).read_text())
        if not isinstance(expect_map, dict):
            sys.exit("--expect-file must be a JSON object: {\"q1\": [globs], \"*\": [globs]}")
    for qi, question in enumerate(_questions(args), 1):
        qid = f"q{qi}"
        globs = list(args.expect) + list(expect_map.get("*", [])) + list(expect_map.get(qid, []))
        expected = expected_from_globs(files, globs)
        cands = retrieve(bm25, question, args.k)
        ctx = assemble_context(cands, chunks, args.context, args.per_file_cap)
        excerpts = [{"n": n, "file": chunks[i].file, "chunk": chunks[i].n, "file_chunks": file_chunk_count[chunks[i].file],
                     "start": chunks[i].start, "end": chunks[i].end, "score": round(s, 3), "text": chunks[i].text}
                    for n, (i, s) in enumerate(ctx, 1)]
        per_file = tier_files(cands, ctx, chunks, file_chunk_count)
        out["questions"].append({
            "id": qid, "question": question, "answer": "",
            "excerpts": excerpts, "files": per_file, "expected": expected,
            "summary": summarize(files, per_file, cands, ctx),
        })
    return out


def _strip_texts(out):
    for q in out["questions"]:
        q["excerpts"] = [{k: v for k, v in e.items() if k != "text"} for e in q["excerpts"]]


def _report(q):
    s = q["summary"]
    print(f"[{q['id']}] retrieved {s['files_retrieved']} files / in context {s['files_in_context']} / "
          f"cited {s['files_cited']}  of {s['files_indexed']} indexed", file=sys.stderr)


def cmd_retrieve(args):
    out = _prepare(args, "external (answer written from excerpts, then `record`)")
    outp = Path(args.out)
    outp.parent.mkdir(parents=True, exist_ok=True)
    for q in out["questions"]:
        ex = outp.parent / f"{q['id']}.excerpts.md"
        ex.write_text(excerpts_markdown(q["question"], q["excerpts"], q["id"]))
        q["excerpts_file"] = str(ex)
        _report(q)
        print(f"      excerpts → {ex}   (write the answer to {outp.parent / (q['id'] + '.answer.md')}, citing [[n]])", file=sys.stderr)
    _strip_texts(out)
    outp.write_text(json.dumps(out, indent=1))
    print(f"wrote {outp}", file=sys.stderr)


def cmd_record(args):
    p = Path(args.coverage)
    out = json.loads(p.read_text())
    q = next((q for q in out["questions"] if q["id"] == args.q), None)
    if q is None:
        sys.exit(f"no question {args.q} in {p}; have {[x['id'] for x in out['questions']]}")
    answer = Path(args.answer).read_text()
    cited = apply_citations(q, answer)
    out["pipeline"]["generator"] = args.generator_label or out["pipeline"].get("generator")
    p.write_text(json.dumps(out, indent=1))
    _report(q)
    for f in sorted(cited):
        print(f"      cited: {f}", file=sys.stderr)
    print(f"updated {p}", file=sys.stderr)


def cmd_run(args):
    label = "stub" if args.generator == "stub" else f"claude:{args.model}"
    out = _prepare(args, label)
    for q in out["questions"]:
        answer = generate_claude(q["question"], q["excerpts"], args.model) if args.generator == "claude" \
            else generate_stub(q["question"], q["excerpts"])
        apply_citations(q, answer)
        _report(q)
    _strip_texts(out)
    outp = Path(args.out)
    outp.parent.mkdir(parents=True, exist_ok=True)
    outp.write_text(json.dumps(out, indent=1))
    print(f"wrote {outp}", file=sys.stderr)


def _common(p):
    p.add_argument("folder", help="folder of documents (the data room export)")
    p.add_argument("-q", action="append", help="question (repeatable)")
    p.add_argument("--questions", help="file with one question per line")
    p.add_argument("-o", "--out", default="coverage.json")
    p.add_argument("--expect", action="append", default=[],
                   help="glob (relative to the folder) of files you expect the answer to draw on; repeatable, applies to every question")
    p.add_argument("--expect-file", help='JSON {"q1": [globs], "q2": [...], "*": [globs for all questions]}')
    p.add_argument("--k", type=int, default=40, help="retrieval candidates (chunks)")
    p.add_argument("--context", type=int, default=12, help="chunks placed in the prompt")
    p.add_argument("--per-file-cap", type=int, default=3, help="max chunks per file in the prompt")
    p.add_argument("--chunk", type=int, default=1200, help="chunk size (chars)")
    p.add_argument("--overlap", type=int, default=150)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    r = sub.add_parser("retrieve", help="index, retrieve, write excerpts for an external answerer")
    _common(r)
    r.set_defaults(func=cmd_retrieve)

    c = sub.add_parser("record", help="parse an answer's [[n]] citations into the cited tier")
    c.add_argument("coverage")
    c.add_argument("--q", required=True, help="question id, e.g. q1")
    c.add_argument("--answer", required=True, help="answer file with [[n]] citations")
    c.add_argument("--generator-label", default="", help="who wrote the answer, for the record (e.g. 'claude-code')")
    c.set_defaults(func=cmd_record)

    u = sub.add_parser("run", help="index, retrieve, answer (stub or API) and record in one go")
    _common(u)
    u.add_argument("--generator", choices=["stub", "claude"], default="stub")
    u.add_argument("--model", default="claude-sonnet-4-5")
    u.set_defaults(func=cmd_run)

    args = ap.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
